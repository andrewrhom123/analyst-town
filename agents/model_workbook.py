"""The PM's workbook as the one model the PM and the agents pass back and forth.

When the PM uploads an edited workbook, the whole file becomes the model, not just the blue input cells:
  - it is stored as uploaded (ModelWorkbook) and is what "download latest" returns;
  - the model-map input cells are found again even after rows, columns or tabs were added (Excel defined
    names, else row label + column header matching) and feed the engine, so the app's valuation stays live;
  - every tab is turned into text (digest) the agents read at their next deep dive, model rebuild or thesis
    update, together with what the PM changed beyond the inputs (structure);
  - the workbook's own valuation (football-field prices, from its formulas) is read back too.
When an agent later updates the model (earnings, new data), its changes are written into the PM's workbook
cell by cell (exports.xlsx_patch: charts, formatting, extra tabs survive), formulas the PM put on input cells
are kept (their values flow back into the engine), and the agent can fill cells outside the inputs (rows the
PM added, period headers) with edit_pm_workbook. The result is the next version of the same workbook.
"""

import copy
import io
import logging
import os
import re
import tempfile
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import select

from agents.financial_model import is_quarterly
from database.db import session_scope
from database.models import ModelWorkbook
from exports import xlsx_patch
from exports.excel_model import MAP_NAME_PREFIX, MAP_SHEET, map_names, map_rows

logger = logging.getLogger(__name__)

MAX_BYTES = 10 * 1024 * 1024
DIGEST_CHARS = 60000  # ~17k tokens: deep dives and rebuilds read the whole workbook
SUMMARY_DIGEST_CHARS = 6000
MAX_STRUCTURE_ITEMS = 60
MAX_AGENT_EDITS = 300
SKIP_SHEETS = {"Historical Tracker"}  # refilled from live SEC actuals at export, not the PM's edits
EXPORT_NOTE = "Model version "  # Summary!B4, rewritten at every export
_PERIOD = re.compile(r"^(\d{4}(?:Q[1-4])?)[AE]$")
_A1 = re.compile(r"(?<![A-Za-z_.\d$])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![\d(A-Za-z_])")


class WorkbookError(ValueError):
    pass


# --- Loading -------------------------------------------------------------------------------

class Book:
    """A workbook opened twice: formulas (f) and the values Excel last calculated (v)."""

    def __init__(self, data: bytes):
        if len(data) > MAX_BYTES:
            raise WorkbookError("File is larger than 10 MB")
        try:
            self.f = load_workbook(io.BytesIO(data), data_only=False)
            self.v = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            raise WorkbookError("That isn't an .xlsx workbook")
        self.data = data
        self._labels: dict[str, dict] = {}
        self._headers: dict[str, dict] = {}

    @property
    def sheets(self) -> list[str]:
        return [n for n in self.f.sheetnames if n != MAP_SHEET]

    def formula(self, sheet: str, ref: str) -> str | None:
        v = self.f[sheet][ref].value
        return v if isinstance(v, str) and v.startswith("=") else None

    def value(self, sheet: str, ref: str, values: dict | None = None):
        """Calculated value: recalculated if given, else Excel's cached one, else the literal."""
        if values is not None and (sheet.upper(), ref) in values:
            return values[(sheet.upper(), ref)]
        v = self.v[sheet][ref].value
        if v is None:
            lit = self.f[sheet][ref].value
            return None if isinstance(lit, str) and lit.startswith("=") else lit
        return v

    # row labels (text in columns A-C) and column headers (nearest text above), for relocating cells
    def row_label(self, sheet: str, r: int) -> str:
        ws = self.f[sheet]
        return " | ".join(v.strip() for c in (1, 2, 3) if isinstance(v := ws.cell(r, c).value, str) and v.strip() and not v.startswith("="))

    def label_rows(self, sheet: str) -> dict[str, list[int]]:
        if sheet not in self._labels:
            out: dict[str, list[int]] = {}
            for r in range(1, self.f[sheet].max_row + 1):
                lab = self.row_label(sheet, r)
                if lab:
                    out.setdefault(lab, []).append(r)
            self._labels[sheet] = out
        return self._labels[sheet]

    def header(self, sheet: str, col: int, r: int) -> str:
        if sheet not in self._headers:
            cols: dict[int, list] = {}
            for row in self.f[sheet].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.strip() and not c.value.startswith("="):
                        cols.setdefault(c.column, []).append((c.row, _norm_header(c.value)))
            self._headers[sheet] = cols
        best = ""
        for rr, text in self._headers[sheet].get(col, []):
            if rr >= r:
                break
            best = text
        return best


def _norm_header(text: str) -> str:
    t = text.strip()
    m = _PERIOD.match(t)
    return m.group(1) if m else t  # 2026Q3E and 2026Q3A are the same column


def read_map(book: Book) -> tuple[dict, list[tuple], dict[str, str]]:
    """(meta, [(sheet, cell, path, kind)], {path: defined name})."""
    import json

    if MAP_SHEET not in book.f.sheetnames:
        raise WorkbookError("This workbook has no hidden _model_map tab (it maps the inputs). Download the latest model, "
                            "make your changes in that file and upload it.")
    mp = book.f[MAP_SHEET]
    try:
        meta = json.loads(mp["A1"].value or "{}")
    except ValueError:
        meta = {}
    entries, names = [], {}
    for row in mp.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        sheet, cell, path, kind = row[:4]
        entries.append((str(sheet), str(cell), str(path), str(kind or "num")))
        if len(row) > 4 and row[4]:
            names[str(path)] = str(row[4])
    return meta, entries, names


def _name_target(book: Book, name: str) -> tuple[str, str] | None:
    dn = book.f.defined_names.get(name)
    if dn is None:
        return None
    try:
        dests = list(dn.destinations)
    except Exception:
        return None
    if len(dests) != 1:
        return None
    sheet, ref = dests[0]
    ref = ref.replace("$", "")
    if sheet not in book.f.sheetnames or not re.fullmatch(r"[A-Z]{1,3}\d+", ref):
        return None  # #REF! (the PM deleted the cell)
    return sheet, ref


# --- Finding mapped cells in an edited workbook -------------------------------------------------

class _Locator:
    """Find where a cell of the reference workbook ended up in the edited one."""

    def __init__(self, ref: Book, up: Book):
        self.ref, self.up = ref, up
        self._anchors: dict[str, list[tuple[int, int]]] = {}

    def anchors(self, sheet: str) -> list[tuple[int, int]]:
        if sheet not in self._anchors:
            a, b = self.ref.label_rows(sheet), self.up.label_rows(sheet)
            self._anchors[sheet] = sorted((a[k][0], b[k][0]) for k in a if len(a[k]) == 1 and len(b.get(k, [])) == 1)
        return self._anchors[sheet]

    def row_shift(self, sheet: str, r: int) -> int:
        shift = 0
        for ref_r, up_r in self.anchors(sheet):
            if ref_r > r:
                break
            shift = up_r - ref_r
        return shift

    def locate(self, sheet: str, ref_cell: str) -> str | None:
        if sheet not in self.up.f.sheetnames:
            return None
        letters, r = xlsx_patch.split_ref(ref_cell)
        c = xlsx_patch.col_index(letters)
        label, head = self.ref.row_label(sheet, r), self.ref.header(sheet, c, r)
        if self.up.row_label(sheet, r) == label and self.up.header(sheet, c, r) == head:
            return ref_cell
        if not label:
            return None
        rows = self.up.label_rows(sheet).get(label, [])
        if not rows:
            return None
        want = r + self.row_shift(sheet, r)
        row = min(rows, key=lambda x: (abs(x - want), x))
        col = c
        if head:
            ws = self.up.f[sheet]
            cands = [k for k in range(1, ws.max_column + 1) if self.up.header(sheet, k, row) == head]
            if cands:
                col = min(cands, key=lambda k: (abs(k - c), k))
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(col)}{row}"


def resolve_map(up: Book, ref: Book, ref_entries: list[tuple], up_names: dict[str, str]) -> tuple[list[tuple], list[str]]:
    """Map entries with their addresses in the edited workbook, and the paths that couldn't be found."""
    loc = _Locator(ref, up)
    resolved, missing = [], []
    for sheet, cell, path, kind in ref_entries:
        target = _name_target(up, up_names[path]) if path in up_names else None
        if target is None:
            found = loc.locate(sheet, cell)
            target = (sheet, found) if found else None
        if target is None:
            missing.append(path)
        else:
            resolved.append((target[0], target[1], path, kind))
    return resolved, missing


# --- What the PM changed beyond the inputs -----------------------------------------------------

def _relative(formula: str, row: int, col: int) -> str:
    def sub(m):
        c, r = xlsx_patch.col_index(m.group(2)), int(m.group(4))
        return f"R[{r - row}]C[{c - col}]"
    return _A1.sub(sub, formula)


def _content(book: Book, sheet: str, cell) -> str | None:
    v = cell.value
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str) and v.startswith("="):
        return "=" + _relative(v[1:], cell.row, cell.column)
    if isinstance(v, float):
        return f"{v:.10g}"
    return str(v)


def _short(v, n: int = 60) -> str:
    s = str(v).replace("\n", " ")
    return s if len(s) <= n else s[:n - 1] + "…"


def structure_changes(ref: Book, up: Book, skip_cells: set[tuple[str, str]]) -> list[str]:
    """Human-readable list of what changed beyond the mapped inputs: tabs, rows, formulas, other numbers."""
    items: list[str] = []
    ref_sheets, up_sheets = set(ref.sheets), set(up.sheets)
    items += [f"Added tab '{s}'" for s in up.sheets if s not in ref_sheets]
    items += [f"Removed tab '{s}'" for s in ref.sheets if s not in up_sheets]
    cell_items, n_cells = [], 0
    for sheet in [s for s in up.sheets if s in ref_sheets and s not in SKIP_SHEETS]:
        a, b = ref.label_rows(sheet), up.label_rows(sheet)
        added = [(rows[0], k) for k, rows in b.items() if k not in a and not k.startswith(EXPORT_NOTE)]
        removed = [(rows[0], k) for k, rows in a.items() if k not in b and not k.startswith(EXPORT_NOTE)]
        items += [f"{sheet}: added row '{_short(k)}' (row {r})" for r, k in sorted(added)]
        items += [f"{sheet}: removed row '{_short(k)}'" for r, k in sorted(removed)]
        loc = _Locator(ref, up)
        ws_up, ws_ref = up.f[sheet], ref.f[sheet]
        width = max(ws_up.max_column, ws_ref.max_column)
        pairs = []  # (row in edited, row in reference, label)
        for k, up_rows in b.items():
            if k not in a:
                continue
            for i, ur in enumerate(up_rows):
                rr = a[k][i] if len(a[k]) == len(up_rows) else min(a[k], key=lambda x: abs(x + loc.row_shift(sheet, x) - ur))
                pairs.append((ur, rr, k))
        anchors = loc.anchors(sheet)
        for ur in range(1, ws_up.max_row + 1):  # rows without a label: align by the nearest labelled row above
            if up.row_label(sheet, ur):
                continue
            shift = next((u - r for r, u in reversed(anchors) if u <= ur), 0)
            if ur - shift >= 1 and not ref.row_label(sheet, ur - shift):
                pairs.append((ur, ur - shift, ""))
        for ur, rr, k in pairs:
            for c in range(4, width + 1):  # labels live in A-C
                cu, cr = ws_up.cell(ur, c), ws_ref.cell(rr, c)
                if (sheet, cu.coordinate) in skip_cells:
                    continue
                new, old = _content(up, sheet, cu), _content(ref, sheet, cr)
                if new == old or cu.value == cr.value or (isinstance(cu.value, str) and cu.value.startswith(EXPORT_NOTE)):
                    continue  # same text = moved, not edited
                n_cells += 1
                if len(cell_items) < MAX_STRUCTURE_ITEMS:
                    head = up.header(sheet, c, ur)
                    what = "formula" if any(isinstance(x, str) and x.startswith("=") for x in (cu.value, cr.value)) else "value"
                    where = ", ".join(x for x in (_short(k, 40), head) if x)
                    cell_items.append(f"{sheet}!{cu.coordinate}{f' ({where})' if where else ''}: {what} "
                                      f"{_short(cr.value, 50) if cr.value is not None else '(empty)'} -> "
                                      f"{_short(cu.value, 50) if cu.value is not None else '(empty)'}")
    out = items[:MAX_STRUCTURE_ITEMS] + cell_items
    if len(items) > MAX_STRUCTURE_ITEMS:
        out.append(f"... and {len(items) - MAX_STRUCTURE_ITEMS} more tab/row changes")
    if n_cells > len(cell_items):
        out.append(f"... and {n_cells - len(cell_items)} more cell changes (see the workbook)")
    return out


# --- Recalculation and the text digest ------------------------------------------------------------

def recalc(data: bytes) -> dict | None:
    """Recalculate every formula ({(SHEET_UPPER, "D8"): value}); None if the workbook can't be computed.
    Needed after agent edits: openpyxl/XML edits leave dependent cells with stale cached values."""
    try:
        import formulas  # heavy; only loaded when a workbook needs computing
        import numpy as np
    except ImportError:
        logger.warning("formulas not installed; workbook values are Excel's cached ones")
        return None
    path = None
    try:
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        sol = formulas.ExcelModel().loads(path).finish().calculate()
        out = {}
        key = re.compile(r"^'\[[^\]]+\](.+)'!([A-Z]{1,3}\d+)$")
        for k, v in sol.items():
            m = key.match(str(k))
            if not m:
                continue
            val = getattr(v, "value", v)
            if isinstance(val, np.ndarray):
                val = val.ravel()[0] if val.size else None
            if isinstance(val, np.generic):
                val = val.item()
            if val is not None and not isinstance(val, (int, float, str, bool, date, datetime)):
                val = str(val)  # XlError (#DIV/0! etc.) or empty
            out[(m.group(1).upper(), m.group(2))] = val
        return out
    except Exception as exc:
        logger.warning("workbook recalculation failed: %s", exc)
        return None
    finally:
        if path:
            try:
                os.remove(path)
            except OSError:
                pass


def _fmt(v, number_format: str) -> str:
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if "%" in (number_format or ""):
            return f"{v * 100:.1f}%"
        a = abs(v)
        if a >= 1000:
            return f"{v:,.0f}"
        if a >= 10:
            return f"{v:,.1f}"
        return f"{v:.4g}"
    if isinstance(v, (date, datetime)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    s = str(v).replace("\r", "").replace("\n", " / ").strip()
    return s if len(s) <= 160 else s[:159] + "…"


DIGEST_LEGEND = ("Every tab, row by row: `r12` = row 12, then cells. `D:450` = a typed-in (hard-coded) value in D12; "
                 "`H=1,840` = a formula in H12 and its calculated value (`=?` not calculated). Percent-formatted cells are shown as %.")


def digest(book: Book, values: dict | None = None, max_chars: int = DIGEST_CHARS, sheets: list[str] | None = None) -> str:
    """The workbook as text, for the agents (values as calculated; formula vs typed-in marked)."""
    parts, total, cut = [], 0, []
    for name in sheets or book.sheets:
        if name not in book.f.sheetnames:
            continue
        ws = book.f[name]
        lines = [f"### {name}"]
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                    continue
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = book.value(name, c.coordinate, values)
                    cells.append(f"{c.column_letter}={_fmt(v, c.number_format) if v not in (None, '') else '?'}")
                else:
                    cells.append(f"{c.column_letter}:{_fmt(c.value, c.number_format)}")
            if cells:
                lines.append(f"r{row[0].row} " + " | ".join(cells))
        block = "\n".join(lines)
        if total + len(block) > max_chars:
            room = max_chars - total
            if room > 2000:
                parts.append(block[:room].rsplit("\n", 1)[0] + "\n… (tab truncated)")
            cut.append(name)
            total = max_chars
            continue
        parts.append(block)
        total += len(block) + 1
    if cut:
        parts.append(f"(Not shown, over the size limit: {', '.join(cut)})")
    return "\n".join(parts)


def needs_recalc(book: Book) -> bool:
    """True when many formulas have no cached value (saved by a tool that doesn't calculate)."""
    n = missing = 0
    for name in book.sheets:
        for row in book.f[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    missing += book.v[name][c.coordinate].value is None
    return n > 0 and missing / n > 0.2


def valuation(book: Book, entries: list[tuple], values: dict | None = None) -> dict:
    out = {}
    for sheet, cell, path, kind in entries:
        if kind == "out" and path.startswith("valuation."):
            v = book.value(sheet, cell, values)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out[path.split(".", 1)[1]] = round(float(v), 4)
    return out


# --- Storage ------------------------------------------------------------------------------------

def get(model_id: int) -> ModelWorkbook | None:
    with session_scope() as s:
        return s.scalar(select(ModelWorkbook).where(ModelWorkbook.model_id == model_id))


def ids_with_workbook(model_ids: list[int]) -> set[int]:
    if not model_ids:
        return set()
    with session_scope() as s:
        return set(s.scalars(select(ModelWorkbook.model_id).where(ModelWorkbook.model_id.in_(model_ids))).all())


def save(model_id: int, data: bytes, entries: list[tuple], meta: dict, digest_text: str, info: dict) -> None:
    """Stamp the map tab with this version's identity/addresses and store the workbook."""
    data = xlsx_patch.patch(data, map_sheet=MAP_SHEET, map_rows=map_rows(meta, entries),
                            names=map_names(entries), name_prefix=MAP_NAME_PREFIX)
    with session_scope() as s:
        s.add(ModelWorkbook(model_id=model_id, data=data, digest=digest_text, info={**info, "map": [list(e) for e in entries]}))


def for_agent(model_id: int) -> dict | None:
    """What the agents read about the PM's workbook for a version."""
    wb = get(model_id)
    if wb is None:
        return None
    info = wb.info or {}
    return {"model_id": model_id, "digest": wb.digest, "structure": info.get("structure") or [],
            "valuation": info.get("valuation") or {}, "missing_inputs": info.get("missing") or [],
            "sheets": info.get("sheets") or []}


# --- Agents working in the PM's workbook ------------------------------------------------------------

EDIT_TOOL = "edit_pm_workbook"
EDIT_TOOL_DEF = {
    "name": EDIT_TOOL,
    "description": ("Write cells in the PM's workbook that your model inputs don't cover: rows, segments or tabs the PM added, "
                    "period headers (e.g. '2026Q3E' -> '2026Q3A' once the quarter is reported), hard-coded figures in the PM's "
                    "own build-outs. Value: a number (percent cells take fractions, 0.25 = 25%), text, or a formula starting "
                    "with '=' for an empty cell. Never overwrites a PM formula; input cells change through "
                    "submit_financial_model. Call it alongside submit_financial_model; later calls add to earlier ones."),
    "input_schema": {
        "type": "object",
        "properties": {"edits": {"type": "array", "maxItems": MAX_AGENT_EDITS, "items": {
            "type": "object",
            "properties": {"sheet": {"type": "string"}, "cell": {"type": "string", "description": "e.g. F18"},
                           "value": {"type": ["number", "string", "null"]}, "reason": {"type": "string"}},
            "required": ["sheet", "cell", "value", "reason"]}}},
        "required": ["edits"],
    },
}

WORKBOOK_GUIDE = """<pm_workbook> is the single model you and the PM share: the PM's own workbook (their edits, added rows, formulas and tabs included), shown tab by tab with calculated values. It is the source of truth.
- Work from it: follow the PM's structure and logic (their segment splits, cost lines, formulas, extra tabs) in your analysis and writing. Where its own valuation differs from the engine's computed model (because the PM's workbook has structure the engine doesn't model), quote the workbook's figures and say why they differ.
- submit_financial_model writes your inputs into its input cells, so keep the same periods (a quarter may flip from E to A), segments and drivers; restructuring is the PM's call. Input cells the PM turned into formulas keep the PM's formula.
- Fill everything else that needs updating (rows the PM added, period headers when a quarter is reported, the PM's own build-outs) with edit_pm_workbook. Never overwrite the PM's formulas.
- Keep every PM assumption unless new data clearly contradicts it, and list each deviation (what, PM value, your value, why) in model_notes."""


def prompt_section(wb: dict, max_chars: int = DIGEST_CHARS) -> str:
    structure = "\n".join(f"- {s}" for s in wb["structure"]) or "- (no changes beyond input cells)"
    missing = f"\nInput cells the app could not find in the workbook (engine kept the prior value): {', '.join(wb['missing_inputs'][:20])}" if wb["missing_inputs"] else ""
    val = f"\nThe workbook's own valuation per share: {wb['valuation']}" if wb["valuation"] else ""
    body = wb["digest"] if len(wb["digest"]) <= max_chars else wb["digest"][:max_chars] + "\n… (truncated)"
    return (f"<pm_workbook>\nWhat the PM changed beyond the input cells (latest upload):\n{structure}{missing}{val}\n\n"
            f"{DIGEST_LEGEND}\n{body}\n</pm_workbook>")


def structure_mismatch(base: dict, new: dict) -> str | None:
    """Why the new inputs can't be written into the PM's workbook (different periods/segments/drivers), or None."""
    if is_quarterly(base) != is_quarterly(new):
        return "the PM's workbook is " + ("quarterly" if is_quarterly(base) else "annual") + "; build the model the same way"
    key = "quarters" if is_quarterly(base) else "fiscal_years"
    strip = lambda ps: [p[:-1] if p[-1:] in "AE" else p for p in ps]  # noqa: E731
    if strip(base.get(key, [])) != strip(new.get(key, [])):
        return f"keep the PM workbook's periods: {base.get(key)} (a period may only flip from E to A)"
    bs, ns = base.get("revenue_segments", []), new.get("revenue_segments", [])
    if [s["name"] for s in bs] != [s["name"] for s in ns]:
        return f"keep the PM workbook's revenue segments, in order: {[s['name'] for s in bs]}"
    for b, n in zip(bs, ns):
        if [d["name"] for d in b["drivers"]] != [d["name"] for d in n["drivers"]]:
            return f"keep the drivers of segment '{b['name']}', in order: {[d['name'] for d in b['drivers']]}"
    return None


class Session:
    """One agent run against the PM's workbook: validates edit_pm_workbook calls and builds the next version."""

    def __init__(self, model_id: int):
        row = get(model_id)
        if row is None:
            raise WorkbookError(f"No workbook for model {model_id}")
        self.model_id = model_id
        self.data = row.data
        self.info = row.info or {}
        self.entries = [tuple(e) for e in self.info.get("map", [])]
        self.book = Book(self.data)
        self.input_cells = {(s, c) for s, c, _p, k in self.entries if k != "out"}
        self.edits: dict[tuple[str, str], tuple] = {}  # (sheet, cell) -> (value, reason)

    def stage(self, edits: list[dict]) -> str:
        errors, ok = [], 0
        if len(self.edits) + len(edits) > MAX_AGENT_EDITS:
            raise WorkbookError(f"At most {MAX_AGENT_EDITS} workbook edits per run")
        for e in edits:
            sheet, cell = str(e.get("sheet", "")), str(e.get("cell", "")).replace("$", "").upper()
            name = next((n for n in self.book.sheets if n.lower() == sheet.lower()), None)
            try:
                xlsx_patch.split_ref(cell)
            except ValueError:
                errors.append(f"{sheet}!{cell}: not a cell reference")
                continue
            if name is None:
                errors.append(f"{sheet}: no such tab (tabs: {', '.join(self.book.sheets)})")
                continue
            if (name, cell) in self.input_cells:
                errors.append(f"{name}!{cell} is a model input; change it through submit_financial_model")
                continue
            if self.book.formula(name, cell):
                errors.append(f"{name}!{cell} holds the PM's formula {self.book.formula(name, cell)}; not overwritten")
                continue
            value = e.get("value")
            if isinstance(value, str) and value.startswith("=") and self.book.f[name][cell].value not in (None, ""):
                errors.append(f"{name}!{cell}: formulas only go into empty cells")
                continue
            self.edits[(name, cell)] = (value, str(e.get("reason", ""))[:300])
            ok += 1
        msg = f"Staged {ok} workbook edit(s); they are written when your model is saved."
        if errors:
            msg += " Not staged: " + "; ".join(errors[:20])
        return msg

    def build(self, base_inputs: dict, new_inputs: dict, comps: dict, symbol: str, labeler) -> dict:
        """Write the new inputs + staged edits into the workbook. Returns {inputs, outputs, data, digest, info,
        changes}; inputs/outputs are re-read from the workbook where the PM's formulas drive input cells."""
        from agents.financial_model import parse_and_compute
        from agents.model_versions import _convert, _get_path, _set_path

        writes, changes, kept = {}, [], []
        for sheet, cell, path, kind in self.entries:
            if kind == "out":
                continue
            try:
                new, old = _get_path(new_inputs, path), _get_path(base_inputs, path)
            except (KeyError, IndexError, TypeError):
                continue
            if new == old or (isinstance(new, (int, float)) and isinstance(old, (int, float)) and abs(new - old) <= 1e-9 * max(1.0, abs(old))):
                continue
            if self.book.formula(sheet, cell):
                kept.append(path)
                continue
            writes[(sheet, cell)] = _to_cell(new, kind)
            changes.append({"path": path, "label": labeler(path, base_inputs), "cell": f"{sheet}!{cell}", "old": old, "new": new})
        for (sheet, cell), (value, reason) in self.edits.items():
            old = self.book.f[sheet][cell].value
            writes[(sheet, cell)] = value
            changes.append({"path": f"cell:{sheet}!{cell}", "label": f"{sheet}!{cell} ({reason})" if reason else f"{sheet}!{cell}",
                            "cell": f"{sheet}!{cell}", "old": old, "new": value})
        data = xlsx_patch.patch(self.data, writes) if writes else self.data
        values = recalc(data) if writes else None
        book = Book(data)
        inputs, outputs = copy.deepcopy(new_inputs), None
        if values is not None:  # PM formulas on input cells drive the engine too
            for sheet, cell, path, kind in self.entries:
                if kind == "out" or not book.formula(sheet, cell):
                    continue
                try:
                    v = _convert(book.value(sheet, cell, values), kind, f"{sheet}!{cell}")
                    if v is not None and v != _get_path(inputs, path):
                        _set_path(inputs, path, v)
                except Exception:
                    continue
        _model, outputs = parse_and_compute(inputs, comps, symbol=symbol)
        info = {k: self.info.get(k) for k in ("structure", "sheets") if self.info.get(k)}
        info.update({"valuation": valuation(book, self.entries, values) if (values is not None or not writes) else {},
                     "missing": self.info.get("missing") or [], "pm_formula_inputs": kept, "recalculated": values is not None})
        text = digest(book, values)
        if writes and values is None:
            text = "(Formula values below are from the PM's last save in Excel; recalculation after the latest edits failed.)\n" + text
        return {"inputs": _model.model_dump(), "outputs": outputs, "data": data, "entries": self.entries,
                "digest": text, "info": info, "changes": changes}


def _to_cell(v, kind: str):
    if v is None:
        return None
    if kind == "pct":
        return float(v) / 100
    if kind == "method":
        return 2 if v == "exit_multiple" else 1
    if kind == "flag":
        return 1 if v else 0
    return v
