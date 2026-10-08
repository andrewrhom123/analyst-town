"""Financial model versions, and the PM's edit round trip.

Every model is a version (oldest = v1): agent deep dives create one, and so does the PM uploading an edited
workbook. Download any version as Excel; edit the blue input cells; upload. The workbook carries a hidden
"_model_map" sheet (which version it came from, and which cell is which model input), so the upload is read
back cell by cell into the model inputs, validated, recomputed by the engine and saved as the new latest
version (source "pm_upload"). From then on the PM's version is what agents use: their files, chat answers,
thesis updates (one is queued to re-anchor the trading levels) and the next deep dive, which starts from the
PM's assumptions and must justify any change to them.
"""

import copy
import io
import json
from datetime import datetime, timezone

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from agents.coverage_files import refresh_files
from agents.financial_model import ModelInputError, ModelInputs, compute_model
from agents.registry import get_context
from database.db import session_scope
from database.models import FinancialModel
from exports.excel_model import MAP_SHEET

MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class ModelUploadError(ValueError):
    pass


def _source(fm: FinancialModel) -> str:
    return (fm.assumptions or {}).get("source") or "agent"


def _aware(dt: datetime) -> datetime:
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)  # SQLite drops tz


def list_versions(symbol: str) -> list[dict]:
    """All model versions for a ticker, newest first."""
    ctx = get_context(symbol)
    if ctx is None:
        return []
    with session_scope() as s:
        rows = s.scalars(select(FinancialModel).where(FinancialModel.ticker_id == ctx.ticker_id).order_by(FinancialModel.date, FinancialModel.id)).all()
        out = []
        for v, fm in enumerate(rows, 1):
            a = fm.assumptions or {}
            o = a.get("outputs") or {}
            dcf = next((m for m in o.get("football_field", []) if m["method"] == "DCF"), {})
            out.append({"model_id": fm.id, "version": v, "date": _aware(fm.date).isoformat(), "source": _source(fm),
                        "note": a.get("note"), "changes": a.get("changes") or [], "parent_model_id": a.get("parent_model_id"),
                        "probability_weighted_price": o.get("probability_weighted_price"),
                        "dcf_price": dcf.get("implied_price"), "latest": v == len(rows)})
        return out[::-1]


def version_meta(symbol: str, model_id: int) -> dict:
    v = next((x for x in list_versions(symbol) if x["model_id"] == model_id), None)
    return {"model_id": model_id, "version": v["version"] if v else None,
            "source_label": "your edits" if v and v["source"] == "pm_upload" else "agent"}


# --- Reading an edited workbook ---------------------------------------------------------

def _set_path(obj, path: str, value) -> None:
    keys = [int(k) if k.isdigit() else k for k in path.split(".")]
    for k in keys[:-1]:
        obj = obj[k]
    obj[keys[-1]] = value


def _get_path(obj, path: str):
    for k in path.split("."):
        obj = obj[int(k)] if k.isdigit() else obj[k]
    return obj


def _convert(raw, kind: str, cell: str):
    if raw is None or raw == "":
        return None
    if isinstance(raw, str):
        try:
            raw = float(raw.replace(",", "").replace("%", "").strip())
        except ValueError:
            raise ModelUploadError(f"{cell}: '{raw}' is not a number")
    if kind == "pct":
        return round(float(raw) * 100, 6)
    if kind == "method":
        return "exit_multiple" if round(float(raw)) == 2 else "perpetuity_growth"
    if kind == "flag":
        return bool(round(float(raw)))
    return float(raw)


def _label(path: str, inputs: dict) -> str:
    years = inputs.get("fiscal_years", [])
    parts = path.split(".")
    yr = lambda i: years[int(i)] if i.isdigit() and int(i) < len(years) else i  # noqa: E731
    names = {"gross_margin_pct": "Gross margin", "adj_ebitda_margin_pct": "Adj. EBITDA margin", "sbc_pct_of_revenue": "SBC % revenue",
             "da_pct_of_revenue": "D&A % revenue", "capex_pct_of_revenue": "Capex % revenue", "cash_tax_rate_pct": "Cash tax rate",
             "nwc_pct_of_revenue_change": "NWC % of revenue change", "wacc_pct": "WACC", "terminal_growth_pct": "Terminal growth",
             "exit_multiple_ev_ebitda": "Exit multiple", "terminal_method": "Terminal method", "sbc_as_cash_cost": "SBC as cash cost",
             "share_price": "Share price", "diluted_shares_mm": "Diluted shares", "total_debt_mm": "Total debt",
             "cash_and_investments_mm": "Cash", "other_ev_adjustments_mm": "Other EV adjustments",
             "base_probability_pct": "Base probability", "probability_pct": "probability",
             "revenue_growth_pct": "revenue growth", "adj_ebitda_margin_pct_s": "EBITDA margin"}
    if parts[0] == "revenue_segments":
        seg = inputs["revenue_segments"][int(parts[1])]
        if parts[2] == "revenue_mm":
            return f"{seg['name']} revenue {yr(parts[3])}"
        if parts[2] == "sotp_ev_to_revenue":
            return f"{seg['name']} SOTP multiple"
        return f"{seg['name']} · {seg['drivers'][int(parts[3])]['name']} {yr(parts[5])}"
    if parts[0] in ("dcf", "capitalization"):
        return f"{'DCF' if parts[0] == 'dcf' else 'Cap'} · {names.get(parts[1], parts[1])}"
    if parts[0] == "scenarios":
        if parts[1] == "base_probability_pct":
            return "Base case probability"
        e_years = [y for y in years if y.endswith("E")]
        field = {"probability_pct": "probability", "revenue_growth_pct": "revenue growth", "adj_ebitda_margin_pct": "EBITDA margin"}[parts[2]]
        return f"{parts[1].title()} {field}" + (f" {e_years[int(parts[3])]}" if len(parts) > 3 and int(parts[3]) < len(e_years) else "")
    return f"{names.get(parts[0], parts[0])}" + (f" {yr(parts[1])}" if len(parts) > 1 else "")


def read_workbook(data: bytes) -> tuple[dict, list[tuple]]:
    """(meta, [(sheet, cell, path, kind, value)]) from an uploaded workbook exported by Analyst Town."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ModelUploadError("File is larger than 10 MB")
    try:
        values = load_workbook(io.BytesIO(data), data_only=True)  # what Excel last calculated
        formulas = load_workbook(io.BytesIO(data), data_only=False)
    except Exception:
        raise ModelUploadError("That isn't an .xlsx workbook")
    if MAP_SHEET not in values.sheetnames:
        raise ModelUploadError("This workbook wasn't exported from Analyst Town (no model map). Download the latest model and edit that.")
    mp = values[MAP_SHEET]
    try:
        meta = json.loads(mp["A1"].value or "{}")
    except ValueError:
        meta = {}
    cells = []
    for sheet, cell, path, kind in mp.iter_rows(min_row=3, values_only=True):
        if not sheet:
            continue
        if sheet not in values.sheetnames:
            raise ModelUploadError(f"Sheet '{sheet}' is missing; keep the tab names as exported")
        raw = values[sheet][cell].value
        if raw is None:  # never calculated (e.g. saved by a tool that doesn't compute formulas): use the literal
            literal = formulas[sheet][cell].value
            raw = None if isinstance(literal, str) and literal.startswith("=") else literal
        cells.append((sheet, cell, path, kind, raw))
    return meta, cells


# --- Applying an upload --------------------------------------------------------------------

def apply_upload(symbol: str, data: bytes, note: str | None = None) -> dict:
    ctx = get_context(symbol)
    if ctx is None:
        raise ModelUploadError(f"{symbol} is not covered")
    meta, cells = read_workbook(data)
    if meta.get("ticker") and meta["ticker"].upper() != ctx.symbol:
        raise ModelUploadError(f"This workbook is for {meta['ticker']}, not {ctx.symbol}")
    with session_scope() as s:
        base = s.get(FinancialModel, meta.get("model_id")) if meta.get("model_id") else None
        if base is None or base.ticker_id != ctx.ticker_id:
            base = s.scalar(select(FinancialModel).where(FinancialModel.ticker_id == ctx.ticker_id).order_by(FinancialModel.date.desc()).limit(1))
        if base is None or not (base.assumptions or {}).get("inputs"):
            raise ModelUploadError(f"No model to edit for {ctx.symbol} yet")
        base_id, research_id = base.id, base.research_id
        base_inputs = copy.deepcopy(base.assumptions["inputs"])
        base_outputs = base.assumptions.get("outputs") or {}
        margin_analysis = base.margin_analysis
    inputs = copy.deepcopy(base_inputs)
    changes = []
    for sheet, cell, path, kind, raw in cells:
        new = _convert(raw, kind, f"{sheet}!{cell}")
        try:
            old = _get_path(base_inputs, path)
        except (KeyError, IndexError, TypeError):
            raise ModelUploadError(f"{sheet}!{cell} maps to {path}, which doesn't exist in the base model")
        if new is None and old is None:
            continue
        if isinstance(old, (int, float)) and isinstance(new, (int, float)) and not isinstance(old, bool):
            if abs(new - old) <= 1e-6 * max(1.0, abs(old)):
                continue
        elif new == old:
            continue
        _set_path(inputs, path, new)
        changes.append({"path": path, "label": _label(path, base_inputs), "cell": f"{sheet}!{cell}", "old": old, "new": new})
    if not changes:
        raise ModelUploadError("No changes found in the blue input cells")
    # models built before an input existed: fill neutral defaults so they still validate
    inputs.setdefault("private_market", {"marks": [], "applied_ev_to_revenue": None,
                                         "rationale": "Not set (model predates private-market inputs)"})
    try:
        model = ModelInputs.model_validate(inputs)
    except ValidationError as exc:
        raise ModelUploadError("Some inputs are invalid: " + "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()[:5]))
    comps = {name: b.get("peers", []) for name, b in ((base_outputs.get("comps") or {}).get("buckets") or {}).items()}
    try:
        outputs = compute_model(model, comps, symbol=ctx.symbol)
    except ModelInputError as exc:
        raise ModelUploadError(str(exc))
    with session_scope() as s:
        fm = FinancialModel(ticker_id=ctx.ticker_id, research_id=research_id, revenue_projections=outputs["rows"],
                            margin_analysis=margin_analysis,
                            assumptions={"inputs": model.model_dump(), "outputs": outputs, "source": "pm_upload",
                                         "note": (note or "").strip()[:1000] or None, "parent_model_id": base_id,
                                         "changes": changes, "uploaded_at": datetime.now(timezone.utc).isoformat()})
        s.add(fm)
        s.flush()
        new_id = fm.id
    refresh_files(ctx.symbol, ["financial_model.json"])
    queued = None
    try:  # re-anchor the trading levels on the PM's model (light model, ~$0.06, runs within research hours)
        from scheduler.jobs import enqueue
        summary = "; ".join(f"{c['label']}: {_fmt(c['old'])} -> {_fmt(c['new'])}" for c in changes[:8])
        queued = enqueue("thesis_update", ctx.symbol, f"PM edited the model: {summary}"[:500], "pm_model")
    except Exception:
        pass
    versions = list_versions(ctx.symbol)
    ff = outputs["football_field"]
    return {"model_id": new_id, "version": versions[0]["version"], "based_on_version": next((v["version"] for v in versions if v["model_id"] == base_id), None),
            "changes": changes, "football_field": [{k: m[k] for k in ("method", "implied_price", "upside_pct")} for m in ff],
            "probability_weighted_price": outputs["probability_weighted_price"], "thesis_update_queued": bool(queued)}


def _fmt(v) -> str:
    return f"{v:,.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else str(v)


def model_for_agent(ticker_id: int) -> dict | None:
    """The latest model as the agent should see it at its next deep dive (inputs + whose version it is)."""
    with session_scope() as s:
        fm = s.scalar(select(FinancialModel).where(FinancialModel.ticker_id == ticker_id).order_by(FinancialModel.date.desc()).limit(1))
        if fm is None or not (fm.assumptions or {}).get("inputs"):
            return None
        a = fm.assumptions
        return {"source": a.get("source") or "agent", "date": _aware(fm.date).date().isoformat(), "note": a.get("note"),
                "pm_changes": [{k: c[k] for k in ("label", "old", "new")} for c in a.get("changes") or []],
                "inputs": a["inputs"]}
